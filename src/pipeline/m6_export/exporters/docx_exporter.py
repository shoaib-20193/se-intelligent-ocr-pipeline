"""
src/pipeline/m6_export/exporters/docx_exporter.py
Exports RenderableDocument to DOCX preserving reading order.
Block-level rendering with native Word tables, headings, and lists.
Bbox coordinates are embedded as small grey metadata text.
"""
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor

from src.pipeline.m6_export.contracts import RenderableDocument


class DocxExporter:
    @staticmethod
    def export(document: RenderableDocument, output_path: str) -> str:
        path = Path(output_path)
        path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()

        for page_idx, page in enumerate(document.pages):
            if page_idx > 0:
                doc.add_page_break()

            sorted_tokens = sorted(page.tokens, key=lambda t: t.reading_order)

            for token in sorted_tokens:
                bbox_info = (
                    f"[bbox: x={token.x:.1f}, y={token.y:.1f}, "
                    f"w={token.width:.1f}, h={token.height:.1f} | "
                    f"col={token.column_id}]"
                )

                b_type = token.block_type.lower()

                if b_type == "heading":
                    heading = doc.add_heading(token.text, level=1)
                    _add_bbox_metadata(heading, bbox_info)

                elif b_type == "table":
                    DocxExporter._draw_table(doc, token)
                    # Add bbox below table
                    para = doc.add_paragraph()
                    _add_bbox_metadata(para, bbox_info)

                elif b_type == "list":
                    for list_line in token.text.split("\n"):
                        clean_text = list_line.lstrip("•-*●○▪► ")
                        if clean_text:
                            para = doc.add_paragraph(clean_text, style="List Bullet")
                    _add_bbox_metadata(para, bbox_info)

                elif b_type == "key_value_pair":
                    para = doc.add_paragraph(token.text)
                    _add_bbox_metadata(para, bbox_info)

                else:
                    # paragraph / unknown
                    para = doc.add_paragraph(token.text)
                    _add_bbox_metadata(para, bbox_info)

        doc.save(str(path))
        return str(path.absolute())

    @staticmethod
    def _draw_table(doc, token):
        """Render table block as a native Word table."""
        table_rows = token.spatial_metadata.get("table_rows")
        
        if table_rows:
            if not table_rows:
                return
            num_cols = max(len(row) for row in table_rows)
            table = doc.add_table(rows=len(table_rows), cols=num_cols)
            table.style = 'Table Grid'

            for row_idx, cells in enumerate(table_rows):
                for col_idx in range(min(num_cols, len(cells))):
                    cell_text = cells[col_idx].strip()
                    cell = table.cell(row_idx, col_idx)
                    cell.text = cell_text
                    
                    if row_idx == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True
        else:
            # Fallback
            rows = token.text.split("\n")
            if not rows:
                return
            first_row_cells = rows[0].split(" | ")
            num_cols = len(first_row_cells)
            table = doc.add_table(rows=len(rows), cols=num_cols)
            table.style = 'Table Grid'

            for row_idx, row_line in enumerate(rows):
                cells = row_line.split(" | ")
                for col_idx in range(min(num_cols, len(cells))):
                    cell_text = cells[col_idx].strip()
                    cell = table.cell(row_idx, col_idx)
                    cell.text = cell_text
                    
                    if row_idx == 0:
                        for paragraph in cell.paragraphs:
                            for run in paragraph.runs:
                                run.font.bold = True


def _add_bbox_metadata(paragraph, bbox_text: str) -> None:
    """
    Appends a small, light-grey run containing bbox coordinates to the
    given paragraph.
    """
    run = paragraph.add_run(f"  {bbox_text}")
    run.font.size = Pt(6)
    run.font.color.rgb = RGBColor(0xCC, 0xCC, 0xCC)
